import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.branch import Branch
from app.models.user import User, UserRole


def _wholesale_user(db: Session, role: UserRole = UserRole.WHOLESALE) -> None:
    branch = Branch(name="Wholesale", phone_number="0", address="Here")
    db.add(branch)
    db.commit()
    db.add(User(id="test-user-id", name="Tester", email="tester@example.com", role=role, branch_id=branch.id))
    db.commit()


def test_export_lists_every_sheet_even_with_no_data(authed_client: TestClient, db_session: Session) -> None:
    _wholesale_user(db_session)

    response = authed_client.get("/api/wholesale/export")

    assert response.status_code == 200
    assert "wholesale-data-" in response.headers["content-disposition"]
    sheets = load_workbook(io.BytesIO(response.content)).sheetnames
    for expected in ("Read Me", "Vouchers", "Voucher Lines", "Orders", "Order Lines", "Stock Now", "Colour Detail"):
        assert expected in sheets


def test_export_carries_the_rows_and_the_csv_zip_matches(authed_client: TestClient, db_session: Session) -> None:
    _wholesale_user(db_session)
    voucher = authed_client.post(
        "/api/wholesale/supplier-vouchers",
        json={
            "supplier_name": "Factory",
            "voucher_date": "2026-09-01",
            "carrier_name": "Cargo",
            "total_packages": 2,
            "lines": [{
                "stock_code": "A1", "description": "Sandal", "product_group": "man",
                "color_breakdown": "black2s,white1s", "unit": "set", "quantity_pairs": 18, "buying_price": 100000,
            }],
        },
    )
    assert voucher.status_code == 201, voucher.text

    workbook = load_workbook(io.BytesIO(authed_client.get("/api/wholesale/export").content))
    header, row = list(workbook["Voucher Lines"].iter_rows(values_only=True))[:2]
    line = dict(zip(header, row))
    assert line["Stock code"] == "A1"
    assert line["Ordered sets"] == 3
    assert line["Line amount (Ks)"] == 300000

    colours = list(workbook["Colour Detail"].iter_rows(values_only=True))[1:]
    assert sorted((r[8], r[9]) for r in colours) == [("black", 12), ("white", 6)]

    archive = zipfile.ZipFile(io.BytesIO(authed_client.get("/api/wholesale/export?format=csv").content))
    assert "02 voucher lines.csv" in archive.namelist()
    assert "A1" in archive.read("02 voucher lines.csv").decode("utf-8-sig")


def test_export_is_wholesale_only(authed_client: TestClient, db_session: Session) -> None:
    _wholesale_user(db_session, UserRole.RETAIL)

    assert authed_client.get("/api/wholesale/export").status_code == 403
